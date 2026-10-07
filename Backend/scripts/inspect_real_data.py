"""Detailed inspection script for real Excel file and Supabase schema."""

import os
import sys
import json
import requests
from dotenv import load_dotenv

# Load env
backend_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
load_dotenv(os.path.join(backend_dir, ".env"))

sys.path.insert(0, backend_dir)

from app.fuel.importer import generate_fuel_preview, detect_header_row, is_summary_or_blank_row, parse_date_value, parse_numeric_value
from app.fuel.mapper import map_headers, normalize_header
import openpyxl

REAL_FILE_PATH = r"C:\Users\HP\Downloads\Transactions_09-02-26_09-09 (1).xlsx"


def inspect_supabase_schema():
    print("==================================================")
    print("1. SUPABASE SCHEMA INSPECTION")
    print("==================================================")
    url = os.getenv("SUPABASE_URL", "").rstrip("/")
    key = os.getenv("SUPABASE_SERVICE_ROLE_KEY") or os.getenv("SUPABASE_ANON_KEY", "")
    print(f"Supabase configured: {bool(url and key)}")
    if not url or not key:
        print("Supabase credentials not configured in environment.")
        return

    headers = {"apikey": key, "Authorization": f"Bearer {key}"}

    # Fetch OpenAPI / schema spec from Supabase
    try:
        res = requests.get(f"{url}/rest/v1/", headers=headers, timeout=10)
        if res.ok:
            schema = res.json()
            definitions = schema.get("definitions", {})
            print("Available tables in PostgREST schema:", list(definitions.keys()))
            
            for tbl_name in ["fuel_imports", "fuel_transactions", "vehicles", "profiles"]:
                if tbl_name in definitions:
                    print(f"\n--- Schema for table: {tbl_name} ---")
                    props = definitions[tbl_name].get("properties", {})
                    for col, details in props.items():
                        print(f"  {col}: {details.get('type')} (format: {details.get('format')}, description: {details.get('description')})")
                else:
                    print(f"\nTable '{tbl_name}' NOT found in definitions.")
        else:
            print(f"Failed to fetch PostgREST OpenAPI spec: {res.status_code} - {res.text[:200]}")
    except Exception as exc:
        print(f"Error querying Supabase schema: {exc}")


def inspect_real_excel():
    print("\n==================================================")
    print("2. REAL EXCEL FILE STRUCTURE & PREVIEW")
    print("==================================================")
    if not os.path.exists(REAL_FILE_PATH):
        print(f"File NOT found at {REAL_FILE_PATH}")
        return

    with open(REAL_FILE_PATH, "rb") as f:
        file_bytes = f.read()

    print(f"File size: {len(file_bytes)} bytes")

    # Load with openpyxl
    wb = openpyxl.load_workbook(io_stream := io.BytesIO(file_bytes), data_only=True)
    print(f"Worksheets in workbook: {wb.sheetnames}")
    
    sheet = wb.active or wb[wb.sheetnames[0]]
    print(f"Active worksheet: '{sheet.title}'")
    print(f"Sheet dimensions: max_row={sheet.max_row}, max_column={sheet.max_column}")

    # Inspect first 10 rows raw
    print("\n--- First 10 Raw Rows ---")
    for r_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        non_empty = [v for v in row if v is not None]
        print(f"Row {r_idx:2d}: ({len(non_empty)} non-empty) -> {row[:10]}")

    # Inspect last 10 rows raw
    print("\n--- Last 10 Raw Rows ---")
    start_last = max(1, sheet.max_row - 10)
    for r_idx, row in enumerate(sheet.iter_rows(min_row=start_last, max_row=sheet.max_row, values_only=True), start=start_last):
        non_empty = [v for v in row if v is not None]
        print(f"Row {r_idx:2d}: ({len(non_empty)} non-empty) -> {row[:10]}")

    # Run preview logic
    print("\n--- Running generate_fuel_preview ---")
    preview = generate_fuel_preview(file_bytes, os.path.basename(REAL_FILE_PATH))
    print(f"Worksheet Name: {preview.worksheet_name}")
    print(f"Detected Header Row: {preview.detected_header_row}")
    print(f"Rows Found: {preview.rows_found}")
    print(f"Rows Candidate: {preview.rows_candidate}")
    print(f"Rows Ignored: {preview.rows_ignored}")
    print(f"Detected Columns ({len(preview.detected_columns)}): {preview.detected_columns}")
    print(f"Suggested Mapping: {json.dumps(preview.suggested_mapping, indent=2)}")
    print(f"Unmapped Columns: {preview.unmapped_columns}")
    print(f"Missing Required Fields: {preview.missing_required_fields}")
    print(f"Validation Warnings ({len(preview.validation_warnings)}):")
    for w in preview.validation_warnings[:10]:
        print(f"  - {w}")
    if len(preview.validation_warnings) > 10:
        print(f"  ... and {len(preview.validation_warnings) - 10} more warnings")

    print("\n--- Sample Parsed Rows (First 5) ---")
    for s in preview.sample_rows:
        print(json.dumps(s, indent=2))


import io

if __name__ == "__main__":
    inspect_supabase_schema()
    inspect_real_excel()
