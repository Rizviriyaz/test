"""
Offline regression tests for Fuel Phase 1: import integrity, wall-clock fingerprints, refunds,
and fail-loud reads. Every outbound HTTP call is blocked by OfflineGuard; Supabase is replaced by
an in-memory fake PostgREST that models how each candidate column type stores and echoes
transaction_date. The real column type is NOT known; these tests prove the dedupe outcome under
each plausible type rather than assuming one.

Run from Backend/:  python scripts/test_fuel_import_integrity.py
"""

import io
import json
import os
import sys
import traceback
from datetime import date, datetime, timedelta, timezone
from unittest.mock import MagicMock, patch

import openpyxl
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
# These tests cover the import with FUEL_IMPORT_REPLACE off (the replacement path has its own suites); a local
# .env that enables it must not change them. load_dotenv never overrides a variable that is already set.
os.environ["FUEL_IMPORT_REPLACE"] = ""

from offline_guard import OfflineGuard  # noqa: E402

from app.main import app  # noqa: E402
from app.fuel import db as fuel_db  # noqa: E402
from app.fuel.analytics import calculate_fuel_summary  # noqa: E402
from app.fuel.anomalies import detect_fuel_waste_and_anomalies, parse_tx_datetime  # noqa: E402
from app.fuel.db import (  # noqa: E402
    FuelDataFetchError,
    build_fingerprint,
    describe_stored_date_shape,
    fetch_existing_transaction_fingerprints,
    fetch_fuel_transactions_for_analytics,
    normalize_date_filter,
)
from app.fuel.importer import FuelImportError, execute_fuel_import, parse_date_value  # noqa: E402
from app.utils.datetime_helpers import parse_tx_wallclock  # noqa: E402

UAE = timezone(timedelta(hours=4))
AUTH = {"Authorization": "Bearer offline-test-token"}
FAKE_USER = {"id": "00000000-0000-0000-0000-000000000001", "email": "t@ats.ae", "role": "admin"}

MAPPING = {
    "transaction_date": "Date",
    "vehicle_identifier": "Vehicle",
    "litres": "Litres",
    "unit_price": "Rate",
    "total_cost": "Amount",
    "invoice_number": "Invoice",
    "odometer_km": "Odo",
}
HEADERS = ["Date", "Vehicle", "Litres", "Rate", "Amount", "Invoice", "Odo"]


# ─── Helpers ───────────────────────────────────────────────────────────────────

def make_xlsx(rows, headers=HEADERS):
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    for r_idx, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class FakeResponse:
    def __init__(self, payload=None, status=200, headers=None):
        self._payload = payload if payload is not None else []
        self.status_code = status
        self.ok = 200 <= status < 300
        self.headers = headers or {}
        self.text = json.dumps(self._payload) if self.ok else "error"

    def json(self):
        return self._payload


class FakePostgrest:
    """
    In-memory fuel_transactions table that mimics PostgreSQL storage + PostgREST rendering of
    transaction_date for a given column type and session TimeZone:
      date         -> stores the calendar date, echoes 'YYYY-MM-DD'
      timestamp    -> stores the naive wall clock, echoes 'YYYY-MM-DDTHH:MM:SS'
      timestamptz  -> interprets the naive input in the session TZ, stores UTC, echoes in session TZ
    Range filters are applied the way PostgreSQL would: literals are interpreted in the same
    session TZ, which is equivalent to comparing wall clocks.
    """

    def __init__(self, column_type="timestamp", session_tz=timezone.utc, reject_negative_litres=False,
                 max_rows=None, send_count=True):
        self.column_type = column_type
        self.session_tz = session_tz
        self.reject_negative_litres = reject_negative_litres
        self.max_rows = max_rows        # server-side cap (PostgREST db-max-rows), may be < requested limit
        self.send_count = send_count    # honour Prefer: count=exact with a Content-Range header
        self.rows = []
        self.get_calls = []
        self.get_headers = []
        self.fail_get_on_call = None  # 1-based index of the GET to fail
        self.fail_mode = "http"

    def serialize(self, written):
        wall = datetime.fromisoformat(written if "T" in written else f"{written}T00:00:00")
        if self.column_type == "date":
            return wall.date().isoformat()
        if self.column_type == "timestamp":
            return wall.strftime("%Y-%m-%dT%H:%M:%S")
        stored_utc = wall.replace(tzinfo=self.session_tz).astimezone(timezone.utc)
        return stored_utc.astimezone(self.session_tz).isoformat()

    def insert(self, transactions, batch_size=100, bearer_token=None, **_kwargs):
        if self.reject_negative_litres and any(t["litres"] < 0 for t in transactions):
            return 0, ["Batch 1 failed: HTTP 400"]
        for t in transactions:
            rec = dict(t)
            rec["id"] = len(self.rows) + 1
            rec["transaction_date"] = self.serialize(t["transaction_date"])
            self.rows.append(rec)
        return len(transactions), []

    def _bounds(self, params):
        lo = hi = None
        if "and" in params:
            inner = params["and"].strip("()").split(",")
            for part in inner:
                _, op, val = part.split(".", 2)
                if op == "gte":
                    lo = val
                elif op == "lte":
                    hi = val
        elif "transaction_date" in params:
            op, val = params["transaction_date"].split(".", 1)
            lo, hi = (val, None) if op == "gte" else (None, val)
        return lo, hi

    def _in_range(self, stored, lo, hi):
        wall = parse_tx_wallclock(stored)[0]
        if self.column_type == "date":
            d = wall.date()
            return (lo is None or d >= datetime.fromisoformat(lo).date()) and (hi is None or d <= datetime.fromisoformat(hi).date())
        return (lo is None or wall >= datetime.fromisoformat(lo)) and (hi is None or wall <= datetime.fromisoformat(hi))

    def get(self, url, params=None, headers=None, timeout=None):
        p = dict(params or {})
        self.get_calls.append(p)
        self.get_headers.append(dict(headers or {}))
        if self.fail_get_on_call == len(self.get_calls):
            if self.fail_mode == "network":
                raise fuel_db.requests.RequestException("simulated network failure")
            return FakeResponse(status=500)
        lo, hi = self._bounds(p)
        rows = [r for r in self.rows if self._in_range(r["transaction_date"], lo, hi)]
        off = int(p.get("offset", 0))
        lim = int(p.get("limit", 1000))
        if self.max_rows is not None:
            lim = min(lim, self.max_rows)
        page = rows[off:off + lim]
        resp_headers = {}
        if self.send_count and "count=exact" in (headers or {}).get("Prefer", ""):
            span = f"{off}-{off + len(page) - 1}" if page else "*"
            resp_headers["content-range"] = f"{span}/{len(rows)}"
        return FakeResponse(page, headers=resp_headers)


def run_import(fake, rows, mapping=MAPPING, headers=HEADERS, **kwargs):
    """Run the real importer + real fingerprint fetch against the fake table."""
    create = MagicMock(side_effect=lambda **kw: 500 + len(fake.rows))
    with patch("app.fuel.importer.fetch_running_imports", return_value=[]), patch("app.fuel.importer.create_fuel_import_record", create), \
         patch("app.fuel.importer.update_fuel_import_record", return_value=True), \
         patch("app.fuel.importer.insert_fuel_transactions_batch", side_effect=fake.insert), \
         patch("app.fuel.importer.resolve_fleet_vehicles", side_effect=AssertionError("must not be called")) as resolve, \
         patch("app.fuel.db._session.get", side_effect=fake.get):
        res = execute_fuel_import(make_xlsx(rows, headers), "statement.xlsx", mapping, FAKE_USER["id"], **kwargs)
    assert not resolve.called, "resolve_fleet_vehicles (Wialon login) must no longer run during import"
    return res, create


COLUMN_TYPES = [
    ("date", timezone.utc),
    ("timestamp", timezone.utc),
    ("timestamptz@UTC", timezone.utc),
    ("timestamptz@Asia/Dubai", UAE),
]


def fake_for(label, tz, **kw):
    return FakePostgrest(column_type=label.split("@")[0], session_tz=tz, **kw)


# ─── 1. Wall-clock normalization ───────────────────────────────────────────────

def test_parse_tx_wallclock_table():
    cases = [
        ("2026-03-01", datetime(2026, 3, 1, 0, 0, 0), "date"),
        ("2026-03-01T10:15:00", datetime(2026, 3, 1, 10, 15, 0), "second"),
        ("2026-03-01 10:15:00", datetime(2026, 3, 1, 10, 15, 0), "second"),
        ("2026-03-01T10:15:00.7478736", datetime(2026, 3, 1, 10, 15, 0), "second"),
        ("2026-03-01T10:15:00+00:00", datetime(2026, 3, 1, 10, 15, 0), "second"),
        ("2026-03-01T10:15:00+04:00", datetime(2026, 3, 1, 10, 15, 0), "second"),  # NOT 06:15
        ("2026-03-01T10:15:00Z", datetime(2026, 3, 1, 10, 15, 0), "second"),
        ("2026-03-01T10:15:00+0400", datetime(2026, 3, 1, 10, 15, 0), "second"),
        ("2026-03-01T00:00:00+00:00", datetime(2026, 3, 1, 0, 0, 0), "date"),
        (datetime(2026, 3, 1, 10, 15, 30, 999), datetime(2026, 3, 1, 10, 15, 30), "second"),
        (datetime(2026, 3, 1, 10, 15, tzinfo=UAE), datetime(2026, 3, 1, 10, 15), "second"),
        (date(2026, 3, 1), datetime(2026, 3, 1), "date"),
    ]
    for value, wall, precision in cases:
        got = parse_tx_wallclock(value)
        assert got == (wall, precision), f"{value!r}: {got}"
    assert parse_tx_wallclock("2026-08-31T10:00:00+04") == (datetime(2026, 8, 31, 10, 0, 0), "second")
    assert describe_stored_date_shape("2026-08-31T10:00:00+04") == "offset:+04"
    for bad in ["garbage", "", None, "2026-02-30", "03/01/2026"]:
        assert parse_tx_wallclock(bad) is None, bad
    print("[PASS] parse_tx_wallclock keeps the printed wall clock and drops offsets without conversion.")


def test_fingerprint_uses_wallclock():
    a = build_fingerprint("DXB 10293", "2026-03-01T10:15:00+04:00", 50, 155, "INV-1", 120000)
    b = build_fingerprint("DXB-10293", "2026-03-01 10:15:00", 50.0, 155.0, "inv-1", 120000.0)
    c = build_fingerprint("DXB 10293", "2026-03-01T06:15:00+00:00", 50, 155, "INV-1", 120000)
    assert a == b, (a, b)
    assert a != c, "an offset must never be converted (06:15Z is a different wall clock)"
    assert "|2026-03-01T10:15:00|" in a
    assert "|2026-03-01T00:00:00|" in build_fingerprint("V", "2026-03-01", 1)
    print("[PASS] build_fingerprint normalizes to wall-clock seconds; offsets ignored, not converted.")


def test_anomaly_reader_agrees_with_fingerprint_reader():
    inputs = [
        "2026-03-01", "2026-03-01T10:15:00", "2026-03-01 10:15:00", "2026-03-01T10:15:00+00:00",
        "2026-03-01T10:15:00+04:00", "2026-03-01T10:15:00Z", "2026-03-01T10:15:00.7478736",
        "2026-03-01T10:15:00+0400",
    ]
    for v in inputs:
        a = parse_tx_datetime(v)
        w = parse_tx_wallclock(v)
        assert a is not None and w is not None, v
        assert a.replace(microsecond=0) == w[0], f"{v}: anomalies={a} wallclock={w[0]}"
    print("[PASS] anomalies.parse_tx_datetime and parse_tx_wallclock agree on the wall clock.")


def test_importer_serialization_unchanged():
    assert parse_date_value(datetime(2026, 3, 1, 10, 15, 30)) == ("2026-03-01T10:15:30", None)
    assert parse_date_value(datetime(2026, 3, 1)) == ("2026-03-01", None)
    assert parse_date_value(46082.4375)[0] == "2026-03-01T10:30:00"  # Excel serial
    assert parse_date_value("2026-03-01 10:15:00.7478736") == ("2026-03-01T10:15:00", None)
    assert parse_date_value("15/03/2026 08:00") == ("2026-03-15T08:00:00", None)
    fake = FakePostgrest("timestamp")
    captured = []
    real_insert = fake.insert
    fake.insert = lambda transactions, **kw: (captured.extend(transactions), real_insert(transactions, **kw))[1]
    run_import(fake, [["2026-03-01 10:15:00", "DXB 10293", 50, 3.1, 155, "INV-1", 1000]])
    assert captured[0]["transaction_date"] == "2026-03-01T10:15:00", captured[0]["transaction_date"]
    for k in ["+", "Z"]:
        assert k not in captured[0]["transaction_date"]
    print("[PASS] Import serialization is unchanged: offset-free wall-clock strings.")


def test_date_filter_bounds_are_wallclock():
    assert normalize_date_filter("2026-08-01") == "2026-08-01T00:00:00"
    assert normalize_date_filter("2026-08-31", is_end=True) == "2026-08-31T23:59:59.999999"
    for label, tz in COLUMN_TYPES:
        fake = fake_for(label, tz)
        fake.insert([{"transaction_date": "2026-08-31T23:30:00", "litres": 10, "vehicle_identifier": "V"}])
        fake.insert([{"transaction_date": "2026-09-01T00:30:00", "litres": 10, "vehicle_identifier": "V"}])
        params = {"and": f"(transaction_date.gte.{normalize_date_filter('2026-08-31')},transaction_date.lte.{normalize_date_filter('2026-08-31', True)})"}
        rows = fake.get("u", params=params).json()
        assert len(rows) == 1 and parse_tx_wallclock(rows[0]["transaction_date"])[0].date() == date(2026, 8, 31), (label, rows)
    print("[PASS] Naive date filters select the same wall-clock day under every simulated column type.")


# ─── 2. Round trip per (unverified) column type ────────────────────────────────

def test_round_trip_dedupe_per_column_type():
    base = [
        ["2026-03-01 08:00:00", "DXB 10293", 50, 3.1, 155, None, None],
        ["2026-03-01 17:00:00", "DXB 10293", 50, 3.1, 155, None, None],  # genuine second refuel
        ["2026-03-02", "SHJ 88211", 40, 3.1, 124, "INV-2", 5000],
    ]
    for label, tz in COLUMN_TYPES:
        fake = fake_for(label, tz)
        first, _ = run_import(fake, base)
        assert first.rows_imported == 3 and first.rows_duplicate == 0, (label, first)

        again, _ = run_import(fake, base)
        assert again.rows_imported == 0 and again.rows_duplicate == 3, (label, again)
        assert again.status == "failed"  # existing behaviour: nothing new imported

        new_day, _ = run_import(fake, [["2026-03-03 10:00:00", "DXB 10293", 50, 3.1, 155, None, None]])
        assert new_day.rows_imported == 1, (label, new_day)

        same_day_new_time, _ = run_import(fake, [["2026-03-01 12:00:00", "DXB 10293", 50, 3.1, 155, None, None]])
        if label == "date":
            # Documented limitation: a date column cannot hold times, so comparison is day-level
            assert same_day_new_time.rows_duplicate == 1, (label, same_day_new_time)
        else:
            assert same_day_new_time.rows_imported == 1, (label, same_day_new_time)
        print(f"[PASS] Round trip under simulated '{label}': re-import fully deduped; "
              f"same-day different-time refuel {'day-level dup (date column)' if label == 'date' else 'kept'}.")


def test_stored_offset_vs_file_time_weak_match():
    fake = FakePostgrest("timestamptz", timezone.utc)
    fake.rows.append({"id": 1, "transaction_date": "2026-03-01T00:00:00+00:00", "vehicle_identifier": "DXB 10293",
                      "litres": 50, "total_cost": 155, "invoice_number": "INV-9", "odometer_km": 1000})
    res, _ = run_import(fake, [["2026-03-01 10:15:00", "DXB 10293", 50, 3.1, 155, "INV-9", 1000]])
    assert res.rows_duplicate == 1 and res.rows_imported == 0, res
    print("[PASS] Date-precision stored row vs timed file row with identical values is a duplicate (weak rule).")


def test_shapes_reported_and_mixed_offsets_warn():
    assert describe_stored_date_shape("2026-03-01") == "date_only"
    assert describe_stored_date_shape("2026-03-01T10:00:00") == "naive"
    assert describe_stored_date_shape("2026-03-01T10:00:00Z") == "offset:+00:00"
    assert describe_stored_date_shape("2026-03-01T10:00:00+04:00") == "offset:+04:00"
    assert describe_stored_date_shape("junk") == "unparseable"
    fake = FakePostgrest("timestamptz", timezone.utc)
    fake.rows += [
        {"id": 1, "transaction_date": "2026-03-01T09:00:00+00:00", "vehicle_identifier": "A", "litres": 10},
        {"id": 2, "transaction_date": "2026-03-01T09:00:00+04:00", "vehicle_identifier": "B", "litres": 10},
    ]
    res, _ = run_import(fake, [["2026-03-01 12:00:00", "C", 20, 3.0, 60, None, None]])
    shapes = res.duplicate_check_scope["stored_value_shapes"]
    assert shapes == {"offset:+00:00": 1, "offset:+04:00": 1}, shapes
    assert any("different UTC offsets" in w for w in res.warnings), res.warnings
    assert "not verified" in res.duplicate_check_scope["timestamp_note"]
    print("[PASS] Stored date shapes reported as evidence; mixed offsets raise a warning.")


# ─── 3. Fingerprint fetch: complete, bounded, fail-loud ────────────────────────

def test_fingerprint_fetch_paginates_with_bounds():
    fake = FakePostgrest("timestamp")
    for i in range(2335):
        fake.rows.append({"id": i + 1, "transaction_date": "2026-08-10T10:00:00", "vehicle_identifier": f"V{i}",
                          "litres": 10, "total_cost": 30})
    with patch("app.fuel.db._session.get", side_effect=fake.get):
        fps = fetch_existing_transaction_fingerprints(date_from="2026-08-01", date_to="2026-08-31")
    assert len(fps) == 2335 and isinstance(fps, set)
    assert [c["offset"] for c in fake.get_calls] == ["0", "1000", "2000"]
    for c in fake.get_calls:
        assert c["order"] == "id.asc" and c["limit"] == "1000"
        assert c["and"] == "(transaction_date.gte.2026-08-01T00:00:00,transaction_date.lte.2026-08-31T23:59:59.999999)"
    assert fps.rows_scanned == 2335
    print("[PASS] Fingerprint fetch pages 1000/1000/335 ordered by id within the date window.")


def test_pagination_not_fooled_by_server_row_cap():
    for send_count in [True, False]:
        fake = FakePostgrest("timestamp", max_rows=500, send_count=send_count)
        fake.rows = [{"id": i + 1, "transaction_date": "2026-08-10T10:00:00", "vehicle_identifier": f"V{i}",
                      "litres": 1, "total_cost": 3} for i in range(1234)]
        with patch("app.fuel.db._session.get", side_effect=fake.get):
            fps = fetch_existing_transaction_fingerprints(date_from="2026-08-01", date_to="2026-08-31")
            rows = fetch_fuel_transactions_for_analytics({})
        assert len(fps) == 1234 and fps.rows_scanned == 1234, (send_count, len(fps))
        assert len(rows) == 1234, (send_count, len(rows))
        assert all(h.get("Prefer") == "count=exact" for h in fake.get_headers)
    print("[PASS] A server max-rows cap below the page size cannot truncate reads (count-based, else empty-page).")


def test_rows_vanishing_mid_read_raises():
    fake = FakePostgrest("timestamp")
    fake.rows = [{"id": i + 1, "transaction_date": "2026-08-10", "vehicle_identifier": "V", "litres": 1} for i in range(1500)]
    real_get = fake.get

    def shrinking_get(url, params=None, headers=None, timeout=None):
        # Server keeps reporting 1500 matching rows but stops returning them after the first page
        resp = real_get(url, params=params, headers=headers, timeout=timeout)
        if len(fake.get_calls) == 1:
            del fake.rows[1000:]
        resp.headers["content-range"] = resp.headers["content-range"].split("/")[0] + "/1500"
        return resp

    with patch("app.fuel.db._session.get", side_effect=shrinking_get):
        try:
            fetch_fuel_transactions_for_analytics({})
            raise AssertionError("expected FuelDataFetchError")
        except FuelDataFetchError as exc:
            assert "1000 of 1500" in str(exc)
    print("[PASS] Fewer rows than the reported count raises instead of returning partial data.")


def test_full_table_reads_use_service_role_when_configured():
    fake = FakePostgrest("timestamp")
    with patch.dict(os.environ, {"SUPABASE_SERVICE_ROLE_KEY": "service-key", "SUPABASE_ANON_KEY": "anon-key"}), \
         patch("app.fuel.db._session.get", side_effect=fake.get):
        fetch_existing_transaction_fingerprints(date_from="2026-08-01", date_to="2026-08-31", bearer_token="user-jwt")
        fetch_fuel_transactions_for_analytics({}, bearer_token="user-jwt")
    assert [h["Authorization"] for h in fake.get_headers] == ["Bearer service-key", "Bearer service-key"]
    fake2 = FakePostgrest("timestamp")
    env = {k: v for k, v in os.environ.items() if k != "SUPABASE_SERVICE_ROLE_KEY"}
    with patch.dict(os.environ, env, clear=True), patch("app.fuel.db._session.get", side_effect=fake2.get):
        fetch_existing_transaction_fingerprints(date_from="2026-08-01", date_to="2026-08-31", bearer_token="user-jwt")
    assert fake2.get_headers[0]["Authorization"] == "Bearer user-jwt"
    print("[PASS] Duplicate-check and analytics reads use the same credentials (service role when configured).")


def test_fingerprint_fetch_raises_on_failure():
    for mode in ["http", "network"]:
        fake = FakePostgrest("timestamp")
        fake.rows = [{"id": i + 1, "transaction_date": "2026-08-10", "vehicle_identifier": "V", "litres": i + 1} for i in range(1500)]
        fake.fail_get_on_call, fake.fail_mode = 2, mode
        with patch("app.fuel.db._session.get", side_effect=fake.get):
            try:
                fetch_existing_transaction_fingerprints(date_from="2026-08-01", date_to="2026-08-31")
                raise AssertionError("expected FuelDataFetchError")
            except FuelDataFetchError:
                pass
    print("[PASS] Fingerprint fetch raises on HTTP and network failure (no silent empty set).")


def test_import_refused_when_duplicate_check_unavailable():
    fake = FakePostgrest("timestamp")
    fake.fail_get_on_call = 1
    inserted = MagicMock(return_value=(0, []))
    create = MagicMock(return_value=1)
    update = MagicMock(return_value=True)
    with patch("app.fuel.importer.fetch_running_imports", return_value=[]), patch("app.fuel.importer.create_fuel_import_record", create), \
         patch("app.fuel.importer.update_fuel_import_record", update), \
         patch("app.fuel.importer.insert_fuel_transactions_batch", inserted), \
         patch("app.fuel.db._session.get", side_effect=fake.get):
        try:
            execute_fuel_import(make_xlsx([["2026-03-05 09:00:00", "AUH 44120", 80, 3.0, 240, None, None]]),
                                "x.xlsx", MAPPING, FAKE_USER["id"])
            raise AssertionError("expected FuelImportError")
        except FuelImportError as exc:
            assert exc.status_code == 503 and "Nothing was imported" in exc.message
    # The record is created (and the import claimed) before the duplicate check, then closed as failed with zero rows
    assert create.called and not inserted.called
    closed = update.call_args.kwargs["updates"]
    assert closed["status"] == "failed" and closed["rows_imported"] == 0 and "Nothing was imported" in closed["error_message"]
    assert fake.get_calls[0]["and"] == "(transaction_date.gte.2026-03-04T00:00:00,transaction_date.lte.2026-03-06T23:59:59.999999)"

    client = TestClient(app)
    fake2 = FakePostgrest("timestamp")
    fake2.fail_get_on_call = 1
    with patch("app.main.authenticate_access_token", return_value=FAKE_USER), \
         patch("app.fuel.importer.fetch_running_imports", return_value=[]), patch("app.fuel.importer.create_fuel_import_record", create), \
         patch("app.fuel.importer.update_fuel_import_record", update), \
         patch("app.fuel.importer.insert_fuel_transactions_batch", inserted), \
         patch("app.fuel.db._session.get", side_effect=fake2.get):
        r = client.post("/api/fuel/import", headers=AUTH, data={"mapping": json.dumps(MAPPING)},
                        files={"file": ("x.xlsx", make_xlsx([["2026-03-05 09:00:00", "AUH 44120", 80, 3.0, 240, None, None]]))})
    assert r.status_code == 503, r.text
    assert not inserted.called and update.call_args.kwargs["updates"]["status"] == "failed"
    print("[PASS] Import refused with 503 and no transaction writes when the duplicate check cannot complete (record closed as failed); window = file dates ±1 day.")


# ─── 4. Refunds, sign rules, warnings, derived values ──────────────────────────

def test_litres_and_price_sign_rules():
    # Rule: litres must be >= 0 (statements carry no refunds); a present price must be numeric and >= 0. Zero is kept.
    fake = FakePostgrest("timestamp")
    res, _ = run_import(fake, [
        ["2026-03-05 09:00:00", "AUH 44120", 80, 3.0, 240, "INV-1", None],    # purchase
        ["2026-03-05 18:00:00", "AUH 44120", -15, None, None, "INV-1R", None],  # negative litres
        ["2026-03-06 09:00:00", "AUH 44120", -10, 3.0, None, None, None],     # negative litres
        ["2026-03-06 10:00:00", "AUH 44120", 15, None, -50, None, None],      # negative cost
        ["2026-03-06 11:00:00", "AUH 44120", 50, 0, None, None, None],        # zero unit price
        ["2026-03-06 12:00:00", "AUH 44120", 0, 3.0, None, None, None],       # zero litres
    ])
    assert res.rows_imported == 3 and res.rows_refund == 0 and res.rows_failed == 3, res
    reasons = [d["reason"] for d in res.failed_row_details]
    assert sum("cannot be negative" in r for r in reasons) == 2, reasons
    assert any("Negative total cost" in r for r in reasons), reasons
    stored = sorted((r["litres"], r["unit_price"], r["total_cost"]) for r in fake.rows)
    assert stored == [(0.0, 3.0, 0.0), (50.0, 0.0, 0.0), (80.0, 3.0, 240.0)], stored
    assert build_fingerprint("V", "2026-03-05", 15, 45) != build_fingerprint("V", "2026-03-05", -15, -45)
    print("[PASS] Negative litres and negative cost rejected with named reasons; zero litres and zero price stored as 0.")


def test_negative_litres_never_reach_the_database():
    fake = FakePostgrest("timestamp", reject_negative_litres=True)
    res, _ = run_import(fake, [
        ["2026-03-05 09:00:00", "AUH 44120", 80, 3.0, 240, None, None],
        ["2026-03-05 18:00:00", "AUH 44120", -15, 3.0, -45, None, None],
    ])
    assert res.rows_imported == 1 and res.rows_refund == 0 and res.rows_failed == 1, res
    assert "cannot be negative" in res.failed_row_details[0]["reason"]
    assert res.status == "completed_with_errors" and "Refund adjustments" not in (res.error_message or "")
    assert [r["litres"] for r in fake.rows] == [80]

    # Failed-row details stay capped at 20
    fake2 = FakePostgrest("timestamp", reject_negative_litres=True)
    bad_rows = [[f"bad-{i}", "AUH 44120", 10, 3.0, 30, None, None] for i in range(25)]
    res2, _ = run_import(fake2, bad_rows + [
        ["2026-03-05 09:00:00", "AUH 44120", 80, 3.0, 240, None, None],
        ["2026-03-05 18:00:00", "AUH 44120", -15, 3.0, -45, None, None],
    ])
    assert res2.rows_failed == 26 and len(res2.failed_row_details) == 20
    print("[PASS] Negative litres are rejected during validation and never sent; purchases kept; details capped at 20.")


def test_negative_net_litres_never_yield_negative_efficiency():
    from app.fuel.efficiency import calculate_fleet_efficiency, calculate_vehicle_efficiency_records
    units = [{"id": 1, "nm": "AUH 44120"}]
    tele = {"AUH 44120": {"mileage_km": 500.0, "engine_hours_seconds": 36000, "has_telemetry": True}}
    txs = [{"vehicle_identifier": "AUH 44120", "litres": -40, "total_cost": -120, "transaction_date": "2026-03-05"}]
    v = calculate_vehicle_efficiency_records(txs, tele, units)[0]
    assert v.fuel_efficiency_l_per_100km is None and v.fuel_cost_per_km is None, v
    assert any("negative net litres" in w for w in v.quality_warnings)
    f = calculate_fleet_efficiency(txs, tele, units)
    assert f.fuel_efficiency_l_per_100km is None and f.fuel_cost_per_km is None, f
    ok = calculate_fleet_efficiency(
        [{"vehicle_identifier": "AUH 44120", "litres": 100, "total_cost": 300},
         {"vehicle_identifier": "AUH 44120", "litres": -20, "total_cost": -60}], tele, units)
    assert ok.fuel_efficiency_l_per_100km == 16.0 and ok.fuel_cost_per_km == 0.48, ok
    print("[PASS] Negative net litres withhold L/100 km and cost/km; net-positive refunds reduce them correctly.")


def test_warnings_and_derived_counts():
    fake = FakePostgrest("timestamp")
    res, _ = run_import(fake, [
        ["03/04/2026", "DXB 10293", 50, 3.1, None, None, None],          # ambiguous date; cost derived
        ["2026-04-05 09:00:00", "DXB 10293", 40, None, 124, None, "12x"],  # unit price derived; bad odometer
        ["2026-04-06 09:00:00", "DXB 10293", 30, None, "AED12.5", None, None],  # unparseable cost -> rejected
    ])
    assert res.rows_imported == 2 and res.rows_failed == 1, res
    assert res.failed_row_details[0]["reason"].startswith("Invalid total cost"), res.failed_row_details
    assert res.rows_date_ambiguous == 1 and any("Ambiguous date" in w for w in res.warnings), res.warnings
    assert res.rows_cost_derived == 1 and res.rows_unit_price_derived == 1
    assert res.rows_optional_parse_failed == 1 and any("could not parse odometer_km" in w for w in res.warnings)
    assert res.file_sha256 and len(res.file_sha256) == 64
    assert res.duplicate_check_scope["date_from"] == "2026-04-02" and res.duplicate_check_scope["date_to"] == "2026-04-06"
    print("[PASS] Ambiguous dates, derived price/cost and unparseable optional cells are reported; an invalid price is rejected.")


def test_in_file_duplicate_details():
    fake = FakePostgrest("timestamp")
    row = ["2026-03-07 09:00:00", "DXB 10293", 50, 3.1, 155, "INV-7", 1000]
    # Both rows carry one identical date: the user confirms they are real refuelling dates
    res, _ = run_import(fake, [row, list(row)], single_date_mode="dated")
    # Approved rule: identical rows of one file are repeated refuels, all imported and reported as a warning
    assert res.rows_imported == 2 and res.rows_duplicate == 0 and res.rows_identical_in_file == 1, res
    assert any("identical to another row" in w for w in res.warnings), res.warnings
    # Re-importing the same file stores nothing new: both copies match the stored rows
    res2, _ = run_import(fake, [row, list(row)], single_date_mode="dated")
    assert res2.rows_imported == 0 and res2.rows_duplicate == 2, res2
    print("[PASS] Identical rows in one file are kept as repeated refuels; a re-import of the file adds nothing.")


# ─── 5. Fail-loud reads ────────────────────────────────────────────────────────

def test_analytics_fetch_complete_or_raise():
    fake = FakePostgrest("timestamp")
    fake.rows = [{"id": i + 1, "transaction_date": "2026-08-10T10:00:00", "vehicle_identifier": "V", "litres": 1,
                  "total_cost": 3} for i in range(1005)]
    with patch("app.fuel.db._session.get", side_effect=fake.get):
        assert len(fetch_fuel_transactions_for_analytics({})) == 1005
    for mode in ["http", "network"]:
        fake.get_calls, fake.fail_get_on_call, fake.fail_mode = [], 2, mode
        with patch("app.fuel.db._session.get", side_effect=fake.get):
            try:
                fetch_fuel_transactions_for_analytics({})
                raise AssertionError("expected FuelDataFetchError")
            except FuelDataFetchError:
                pass
    print("[PASS] Analytics fetch returns every page, or raises on a mid-way failure (no partial data).")


def test_routes_return_503_instead_of_partial_kpis():
    client = TestClient(app)
    with patch("app.main.authenticate_access_token", return_value=FAKE_USER), \
         patch("app.fuel.db._session.get", return_value=FakeResponse(status=500)):
        r = client.get("/api/fuel/analytics/summary", headers=AUTH)
        assert r.status_code == 503 and "withheld" in r.json()["detail"], r.text
        with patch("app.fuel.routes.login", return_value="sid"), \
             patch("app.fuel.routes.get_all_vehicles", return_value=[{"id": 1, "nm": "DXB 10293"}]):
            r2 = client.get("/api/fuel/analytics/vehicle-detail?vehicle_identifier=DXB 10293", headers=AUTH)
            assert r2.status_code == 503, r2.text
    print("[PASS] /analytics/summary and /analytics/vehicle-detail return 503 when Supabase reads fail.")


def test_transactions_route_labels_refunds():
    client = TestClient(app)
    rows = [
        {"id": 1, "import_id": 1, "transaction_date": "2026-03-05T09:00:00", "vehicle_identifier": "AUH 44120", "litres": 80.0, "total_cost": 240.0},
        {"id": 2, "import_id": 1, "transaction_date": "2026-03-05T18:00:00", "vehicle_identifier": "AUH 44120", "litres": -15.0, "total_cost": -45.0},
    ]
    with patch("app.main.authenticate_access_token", return_value=FAKE_USER), \
         patch("app.fuel.routes.fetch_fuel_transactions_for_analytics", return_value=rows), \
         patch("app.fuel.routes.login", return_value="sid"), \
         patch("app.fuel.routes.get_all_vehicles", return_value=[{"id": 1, "nm": "AUH 44120"}]):
        r = client.get("/api/fuel/transactions", headers=AUTH)
    assert r.status_code == 200, r.text
    kinds = {t["id"]: t["transaction_kind"] for t in r.json()["transactions"]}
    assert kinds == {1: "purchase", 2: "refund_adjustment"}, kinds
    print("[PASS] /api/fuel/transactions labels negative-litre rows as refund_adjustment.")


# ─── 6. Read-side refund handling ──────────────────────────────────────────────

def test_summary_net_vs_gross():
    s = calculate_fuel_summary([
        {"vehicle_identifier": "A", "litres": 100, "total_cost": 300, "fuel_type": "DIESEL"},
        {"vehicle_identifier": "A", "litres": 50, "total_cost": 160, "fuel_type": "DIESEL"},
        {"vehicle_identifier": "A", "litres": -20, "total_cost": -60, "fuel_type": "DIESEL"},
    ])
    assert s.total_litres == 130.0 and s.total_cost == 400.0
    assert s.gross_purchased_litres == 150.0 and s.gross_purchased_cost == 460.0 and s.purchase_count == 2
    assert s.refund_count == 1 and s.refund_litres == -20.0 and s.refund_cost == -60.0
    assert s.average_unit_price == round(460 / 150, 2)
    assert s.transaction_count == 3
    plain = calculate_fuel_summary([{"vehicle_identifier": "A", "litres": 100, "total_cost": 310}])
    assert plain.average_unit_price == 3.1 and plain.refund_count == 0
    print("[PASS] Summary totals are net of refunds; average price uses purchases only; unchanged without refunds.")


def test_anomalies_skip_refunds():
    txs = [
        {"id": 1, "vehicle_identifier": "A", "transaction_date": "2026-08-01T10:00:00", "litres": 60, "total_cost": 180, "odometer_km": 1000},
        {"id": 2, "vehicle_identifier": "A", "transaction_date": "2026-08-01T20:00:00", "litres": -20, "total_cost": -60, "odometer_km": 50},
        {"id": 3, "vehicle_identifier": "A", "transaction_date": "2026-08-02T08:00:00", "litres": 55, "total_cost": 165, "odometer_km": 1200},
    ]
    res = detect_fuel_waste_and_anomalies(txs, [], 5.0, 2.0)
    assert all(a.transaction_id != 2 for a in res.anomalies), [a.anomaly_type for a in res.anomalies]
    rapid = [a for a in res.anomalies if a.anomaly_type == "rapid_refuel"]
    assert len(rapid) == 1 and rapid[0].transaction_id == 3 and rapid[0].details["previous_litres"] == 60.0
    assert not [a for a in res.anomalies if a.anomaly_type == "odometer_inconsistency"]
    print("[PASS] Refund rows never trigger refuel rules; interval/odometer reference stays the previous purchase.")


# ─── Runner ────────────────────────────────────────────────────────────────────

TESTS = [
    test_parse_tx_wallclock_table,
    test_fingerprint_uses_wallclock,
    test_anomaly_reader_agrees_with_fingerprint_reader,
    test_importer_serialization_unchanged,
    test_date_filter_bounds_are_wallclock,
    test_round_trip_dedupe_per_column_type,
    test_stored_offset_vs_file_time_weak_match,
    test_shapes_reported_and_mixed_offsets_warn,
    test_fingerprint_fetch_paginates_with_bounds,
    test_pagination_not_fooled_by_server_row_cap,
    test_rows_vanishing_mid_read_raises,
    test_full_table_reads_use_service_role_when_configured,
    test_fingerprint_fetch_raises_on_failure,
    test_import_refused_when_duplicate_check_unavailable,
    test_litres_and_price_sign_rules,
    test_negative_litres_never_reach_the_database,
    test_negative_net_litres_never_yield_negative_efficiency,
    test_warnings_and_derived_counts,
    test_in_file_duplicate_details,
    test_analytics_fetch_complete_or_raise,
    test_routes_return_503_instead_of_partial_kpis,
    test_transactions_route_labels_refunds,
    test_summary_net_vs_gross,
    test_anomalies_skip_refunds,
]


if __name__ == "__main__":
    os.environ.setdefault("SUPABASE_URL", "https://offline.invalid")
    print("=" * 60)
    print("FUEL PHASE 1 - IMPORT INTEGRITY (OFFLINE)")
    print("=" * 60)
    guard = OfflineGuard().start()
    failures = 0
    try:
        for t in TESTS:
            try:
                t()
            except Exception:
                failures += 1
                print(f"[FAIL] {t.__name__}")
                traceback.print_exc()
    finally:
        guard.stop()
    guard.assert_no_live_calls()
    print("=" * 60)
    print(f"{len(TESTS) - failures}/{len(TESTS)} test functions passed")
    sys.exit(1 if failures else 0)
