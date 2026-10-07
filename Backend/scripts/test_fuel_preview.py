"""Test suite for Fuel Import Preview engine and FastAPI routes."""

import io
import sys
import os
from datetime import datetime
import openpyxl
from fastapi.testclient import TestClient

# Add parent directory to path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.main import app
from app.fuel.importer import generate_fuel_preview, parse_date_value, parse_numeric_value
from app.fuel.mapper import map_headers, normalize_header


def create_test_excel(headers, rows, sheet_name="Fuel Report", start_row=1, title_rows=None):
    """Helper to generate an in-memory .xlsx file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    current_row = 1
    if title_rows:
        for t in title_rows:
            ws.cell(row=current_row, column=1, value=t)
            current_row += 1

    while current_row < start_row:
        current_row += 1

    # Write headers
    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=current_row, column=col_idx, value=h)
    current_row += 1

    # Write rows
    for r in rows:
        for col_idx, val in enumerate(r, start=1):
            ws.cell(row=current_row, column=col_idx, value=val)
        current_row += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_tests():
    print("==================================================")
    print("RUNNING FUEL IMPORT PREVIEW TEST SUITE")
    print("==================================================")

    # Test 1: Normalizer & Mapper
    print("\n--- Test 1: Header Normalization & Canonical Mapping ---")
    headers_sample = [
        "Txn Date", "Plate No", "Fuel Type", "Qty (Ltr)", "Unit Rate",
        "Total (AED)", "Odometer Reading", "Station Name", "Invoice #",
        "Fuel Card No", "Driver Name", "Ref No", "Remarks", "Extra Unknown Col"
    ]
    mapping, details, unmapped, missing_req = map_headers(headers_sample)
    print(f"Detected mapping: {mapping}")
    print(f"Unmapped columns: {unmapped}")
    print(f"Missing required: {missing_req}")
    assert mapping["transaction_date"] == "Txn Date"
    assert mapping["vehicle_identifier"] == "Plate No"
    assert mapping["litres"] == "Qty (Ltr)"
    assert mapping["unit_price"] == "Unit Rate"
    assert mapping["total_cost"] == "Total (AED)"
    assert mapping["odometer_km"] == "Odometer Reading"
    assert mapping["station_name"] == "Station Name"
    assert mapping["invoice_number"] == "Invoice #"
    assert mapping["card_number"] == "Fuel Card No"
    assert mapping["driver_name"] == "Driver Name"
    assert mapping["reference_number"] == "Ref No"
    assert mapping["notes"] == "Remarks"
    assert "Extra Unknown Col" in unmapped
    assert len(missing_req) == 0
    print("[PASS] Header mapping works accurately.")

    # Test 2: Number & Date Parsing
    print("\n--- Test 2: Value Parsers ---")
    num1, _ = parse_numeric_value("AED 250.75")
    num2, _ = parse_numeric_value("1,450.50 Ltr")
    num3, _ = parse_numeric_value(85.2)
    assert num1 == 250.75
    assert num2 == 1450.50
    assert num3 == 85.2

    d1, _ = parse_date_value("2026-03-15")
    d2, _ = parse_date_value("15/03/2026")
    d3, _ = parse_date_value(datetime(2026, 3, 15, 10, 30))
    assert d1 == "2026-03-15"
    assert d2 == "2026-03-15"
    assert "2026-03-15" in d3
    print("[PASS] Value parsing works accurately.")

    # Test 3: Full In-Memory Excel Preview (ADNOC style with Title & Summary rows)
    print("\n--- Test 3: Complex Excel Preview with Header Detection & Ignored Rows ---")
    excel_bytes = create_test_excel(
        headers=["Date", "Vehicle No", "Product", "Volume", "Price", "Amount", "Odo KM", "Station", "Invoice No"],
        rows=[
            ["2026-03-01", "DXB-10293", "Diesel", "75.50 AED", "3.10", "234.05", "125000", "ADNOC 842", "INV-1001"],
            ["2026-03-02", "SHJ-88211", "Special 95", "45.00", "3.02", "135.90", "89200", "ADNOC 110", "INV-1002"],
            ["2026-03-03", "AUH-44120", "Diesel", "120.00", "3.10", "372.00", "210400", "ADNOC 503", "INV-1003"],
            [None, None, None, None, None, None, None, None, None],  # Blank row
            ["Grand Total", "", "", "240.50", "", "741.95", "", "", ""]  # Total row
        ],
        title_rows=["AL THIKRA TRANSPORT SERVICES", "MONTHLY FUEL STATEMENT - MARCH 2026"],
        start_row=4
    )

    preview = generate_fuel_preview(excel_bytes, "adnoc_march_2026.xlsx")
    print(f"Worksheet: {preview.worksheet_name}")
    print(f"Detected Header Row: {preview.detected_header_row}")
    print(f"Rows found: {preview.rows_found}")
    print(f"Rows candidate: {preview.rows_candidate}")
    print(f"Rows ignored: {preview.rows_ignored}")
    print(f"Sample rows count: {len(preview.sample_rows)}")
    assert preview.detected_header_row == 4
    assert preview.rows_candidate == 3
    assert preview.rows_ignored == 2  # 1 blank + 1 grand total
    assert len(preview.sample_rows) == 3
    assert preview.sample_rows[0]["vehicle_identifier"] == "DXB-10293"
    assert preview.sample_rows[0]["litres"] == 75.50
    assert preview.sample_rows[0]["station_name"] == "ADNOC 842"
    print("[PASS] Excel preview correctly detected headers at row 4 and filtered total/blank rows.")

    # Test 4: Missing Required Fields Detection
    print("\n--- Test 4: Missing Required Columns Detection ---")
    incomplete_excel = create_test_excel(
        headers=["Date", "Station", "Amount"],
        rows=[["2026-03-01", "Station 1", "100.0"]]
    )
    incomplete_preview = generate_fuel_preview(incomplete_excel, "incomplete.xlsx")
    print(f"Missing required fields: {incomplete_preview.missing_required_fields}")
    assert "vehicle_identifier" in incomplete_preview.missing_required_fields
    assert "litres" in incomplete_preview.missing_required_fields
    print("[PASS] Missing required fields properly flagged.")

    # Test 5: FastAPI TestClient integration
    print("\n--- Test 5: FastAPI Route Integration & Auth Check ---")
    client = TestClient(app)

    # Health check
    res_health = client.get("/api/health")
    print(f"/api/health status: {res_health.status_code}")
    assert res_health.status_code == 200

    # Unauthenticated request to /api/fuel/import/preview should be rejected with 401
    res_unauth = client.post(
        "/api/fuel/import/preview",
        files={"file": ("test.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    print(f"Unauthenticated /api/fuel/import/preview status: {res_unauth.status_code}")
    assert res_unauth.status_code == 401

    # Mock auth for testing endpoint execution directly
    from unittest.mock import patch
    with patch("app.main.authenticate_access_token", return_value={"id": "test-uuid-123", "email": "admin@ats.ae", "role": "admin"}):
        auth_headers = {"Authorization": "Bearer test-valid-token"}

        # 5a: Invalid extension (.pdf)
        res_bad_ext = client.post(
            "/api/fuel/import/preview",
            headers=auth_headers,
            files={"file": ("report.pdf", b"%PDF-1.4...", "application/pdf")}
        )
        print(f"Invalid extension status: {res_bad_ext.status_code}")
        assert res_bad_ext.status_code == 400
        assert "Unsupported file format" in res_bad_ext.json()["detail"]

        # 5b: Corrupted Excel
        res_corrupt = client.post(
            "/api/fuel/import/preview",
            headers=auth_headers,
            files={"file": ("broken.xlsx", b"not an excel zip", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        print(f"Corrupt file status: {res_corrupt.status_code}")
        assert res_corrupt.status_code == 422

        # 5c: Successful preview response
        res_ok = client.post(
            "/api/fuel/import/preview",
            headers=auth_headers,
            files={"file": ("adnoc_march_2026.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        print(f"Authenticated preview status: {res_ok.status_code}")
        assert res_ok.status_code == 200
        data = res_ok.json()
        print("Preview response keys:", list(data.keys()))
        assert data["file_name"] == "adnoc_march_2026.xlsx"
        assert data["worksheet_name"] == "Fuel Report"
        assert data["detected_header_row"] == 4
        assert data["rows_candidate"] == 3
        assert data["suggested_mapping"]["vehicle_identifier"] == "Vehicle No"
        assert data["suggested_mapping"]["litres"] == "Volume"
        assert len(data["sample_rows"]) == 3

    print("\n==================================================")
    print("ALL TESTS PASSED SUCCESSFULLY!")
    print("==================================================")


if __name__ == "__main__":
    run_tests()
